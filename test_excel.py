"""Тесты модулей Excel и CSV (библиотека/excel.dan, библиотека/csv.dan).

Книги, созданные Даном, читаются openpyxl и LibreOffice; книги openpyxl и
LibreOffice (общие строки, даты, формулы, логические значения, пропуски)
читаются Даном. CSV сверяется с модулем csv языка Python на случайных
таблицах.
"""

import csv
import datetime
import importlib.util
import io
import os
import random
import shutil
import subprocess
import tempfile
import unittest

from test_kodogen import КОМПИЛЯТОР_ЕСТЬ
from test_word import выполнить

ЕСТЬ_OPENPYXL = importlib.util.find_spec("openpyxl") is not None
ЕСТЬ_OFFICE = shutil.which("soffice") is not None

СОЗДАНИЕ = r'''подключить "excel"
если ошибка о = excel.создать("книга.xlsx", [
    excel.Лист("Октябрь", [["Товар", "Количество", "Цена", "Дата", "Код"],
                          ["Книга «Дан» & <ко>", "2", "350.50", "26.10.2026", "007"],
                          ["Тетрадь", "-10", "0", "31.12.2026", "+7 930"],
                          ["Длинное число", "1234567890123456", "1.", ".5", "1e5"]]),
    excel.Лист("Второй", [["а"], [], ["", "б"]], шапка = ложь)
]) {
    вывод(о)
}
'''


@unittest.skipUnless(КОМПИЛЯТОР_ЕСТЬ, "нет компилятора C")
class ТестыExcel(unittest.TestCase):

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.п = self.tmp.name

    def tearDown(self):
        self.tmp.cleanup()

    def файл(self, имя):
        return os.path.join(self.п, имя)

    def test_созданная_книга_читается_даном(self):
        код, вывод, ош = выполнить(self.п, СОЗДАНИЕ + '''если пусть к = excel.открыть("книга.xlsx") {
    вывод(к.листы())
    для с в к.лист("Октябрь") иначе [] { вывод(с) }
    вывод(к.лист_по_номеру(1) иначе [])
    если ошибка о = к.лист("Нет") { вывод(о) }
}
вывод(excel.столбец(0) + " " + excel.столбец(25) + " " + excel.столбец(26) + " " + excel.столбец(701) + " " + excel.столбец(702))
''')
        self.assertEqual(код, 0, ош)
        self.assertEqual(вывод.splitlines(), [
            '["Октябрь", "Второй"]',
            '["Товар", "Количество", "Цена", "Дата", "Код"]',
            '["Книга «Дан» & <ко>", "2", "350.50", "26.10.2026", "007"]',
            '["Тетрадь", "-10", "0", "31.12.2026", "+7 930"]',
            '["Длинное число", "1234567890123456", "1.", ".5", "1e5"]',
            '[["а", ""], ["", ""], ["", "б"]]',
            "в книге нет листа «Нет»; листы: Октябрь, Второй",
            "A Z AA ZZ AAA"])

    @unittest.skipUnless(ЕСТЬ_OPENPYXL, "нет openpyxl (pip install openpyxl)")
    def test_созданная_книга_читается_openpyxl(self):
        import openpyxl
        код, _, ош = выполнить(self.п, СОЗДАНИЕ)
        self.assertEqual(код, 0, ош)
        к = openpyxl.load_workbook(self.файл("книга.xlsx"))
        self.assertEqual(к.sheetnames, ["Октябрь", "Второй"])
        л = к["Октябрь"]
        self.assertEqual([[я.value for я in с] for с in л.iter_rows()], [
            ["Товар", "Количество", "Цена", "Дата", "Код"],
            ["Книга «Дан» & <ко>", 2, 350.5, datetime.datetime(2026, 10, 26), "007"],
            ["Тетрадь", -10, 0, datetime.datetime(2026, 12, 31), "+7 930"],
            # число длиннее 15 цифр и необычные записи остаются текстом
            ["Длинное число", "1234567890123456", "1.", ".5", "1e5"]])
        self.assertTrue(л["A1"].font.b)
        self.assertFalse(л["A2"].font.b)
        self.assertEqual(л["D2"].number_format, "dd\\.mm\\.yyyy")
        self.assertEqual(л.freeze_panes, "A2")
        self.assertIsNone(к["Второй"].freeze_panes)

    @unittest.skipUnless(ЕСТЬ_OPENPYXL, "нет openpyxl (pip install openpyxl)")
    def test_чтение_книги_openpyxl(self):
        import openpyxl
        к = openpyxl.Workbook()
        л = к.active
        л.title = "Данные"
        л["A1"] = "Имя"
        л["B1"] = "Сумма"
        л["C1"] = "Дата"
        л["D1"] = "Ответ"
        л["A2"] = "Иванов"
        л["B2"] = 1500.25
        л["C2"] = datetime.date(2026, 10, 26)
        л["D2"] = True
        л["A3"] = "Петров"
        л["B3"] = 42
        л["C3"] = datetime.datetime(2026, 1, 5, 14, 30)
        л["D3"] = False
        л["F5"] = "далеко"          # пропуски строк и столбцов
        л["B6"] = "=B2+B3"          # формула без вычисленного значения
        второй = к.create_sheet("Пустой")
        к.save(self.файл("из_openpyxl.xlsx"))
        код, вывод, ош = выполнить(self.п, '''подключить "excel"
если пусть к = excel.открыть("из_openpyxl.xlsx") {
    вывод(к.листы())
    для с в к.лист("Данные") иначе [] { вывод(с) }
    вывод(к.лист("Пустой") иначе [["ошибка"]])
} иначе ошибка о {
    вывод(о)
}
''')
        self.assertEqual(код, 0, ош)
        self.assertEqual(вывод.splitlines(), [
            '["Данные", "Пустой"]',
            '["Имя", "Сумма", "Дата", "Ответ", "", ""]',
            '["Иванов", "1500.25", "26.10.2026", "истина", "", ""]',
            '["Петров", "42", "05.01.2026 14:30", "ложь", "", ""]',
            '["", "", "", "", "", ""]',
            '["", "", "", "", "", "далеко"]',
            '["", "=B2+B3", "", "", "", ""]',
            "[]"])

    @unittest.skipUnless(ЕСТЬ_OFFICE, "нет LibreOffice")
    def test_книги_и_офис(self):
        # книга Дана открывается офисом, а книга офиса (общие строки, свои стили) - Даном
        код, _, ош = выполнить(self.п, СОЗДАНИЕ)
        self.assertEqual(код, 0, ош)
        р = subprocess.run(["soffice", "--headless", "--convert-to", "csv:Text - txt - csv (StarCalc):44,34,76",
                            "--outdir", self.п, self.файл("книга.xlsx")], capture_output=True, text=True, timeout=180)
        self.assertEqual(р.returncode, 0, р.stderr)
        with open(self.файл("книга.csv"), encoding="utf-8") as ф:
            строки = list(csv.reader(ф))
        self.assertEqual(строки[1][0], "Книга «Дан» & <ко>")
        self.assertEqual(строки[1][1], "2")
        self.assertEqual(строки[1][3], "26.10.2026")
        with open(self.файл("исходная.csv"), "w", encoding="utf-8") as ф:
            ф.write("Имя,Сумма\nИванов,1500.25\nПетров,42\n")
        р = subprocess.run(["soffice", "--headless", "--infilter=CSV:44,34,76", "--convert-to", "xlsx",
                            "--outdir", self.п, self.файл("исходная.csv")], capture_output=True, text=True, timeout=180)
        self.assertEqual(р.returncode, 0, р.stderr)
        код, вывод, ош = выполнить(self.п, '''подключить "excel"
если пусть к = excel.открыть("исходная.xlsx") {
    для с в к.лист_по_номеру(0) иначе [] { вывод(с) }
}
''')
        self.assertEqual(код, 0, ош)
        self.assertEqual(вывод.splitlines(), ['["Имя", "Сумма"]', '["Иванов", "1500.25"]', '["Петров", "42"]'])

    def test_ошибки(self):
        with open(self.файл("текст.xlsx"), "w", encoding="utf-8") as ф:
            ф.write("это не книга")
        код, вывод, ош = выполнить(self.п, '''подключить "excel"
если ошибка о = excel.открыть("нет.xlsx") { вывод(о) }
если ошибка о = excel.открыть("текст.xlsx") { вывод(о) }
если ошибка о = excel.создать("к.xlsx", []) { вывод(о) }
если ошибка о = excel.создать("к.xlsx", [excel.Лист("а/б", [])]) { вывод(о) }
если ошибка о = excel.создать("к.xlsx", [excel.Лист("Лист", []), excel.Лист("лист", [])]) { вывод(о) }
''')
        self.assertEqual(код, 0, ош)
        self.assertEqual(вывод.splitlines(), [
            "файл «нет.xlsx» не найден",
            "это не книга Excel в формате .xlsx (старый формат .xls не поддерживается: сохраните её в Excel как .xlsx)",
            "в книге Excel должен быть хотя бы один лист",
            "имя листа «а/б» не может содержать знак /",
            "два листа с именем «лист»"])


@unittest.skipUnless(КОМПИЛЯТОР_ЕСТЬ, "нет компилятора C")
class ТестыCSV(unittest.TestCase):

    def test_случайные_таблицы_как_python(self):
        сл = random.Random(26)
        знаки = 'аб c,;"\n\r\t1«»'
        таблицы = []
        for _ in range(30):
            ширина = сл.randint(1, 4)
            таблицы.append([["".join(сл.choice(знаки) for _ in range(сл.randint(0, 6))) for _ in range(ширина)]
                            for _ in range(сл.randint(1, 5))])
        with tempfile.TemporaryDirectory() as п:
            for н, т in enumerate(таблицы):
                for разделитель, имя in ((",", "з"), (";", "т")):
                    буфер = io.StringIO()
                    csv.writer(буфер, delimiter=разделитель).writerows(т)
                    with open(os.path.join(п, f"{имя}{н}.csv"), "w", encoding="utf-8", newline="") as ф:
                        ф.write(буфер.getvalue())
            программа = '''подключить "csv"
для н в 0..29 {
    // файл читается байтами: прочитать() заменяет переводы строк CR LF на LF, а они бывают в значениях
    пусть т = csv.разобрать_с_разделителем(стр(прочитать_байты("з" + стр(н) + ".csv")), ",") иначе [["ошибка"]]
    записать("итог_з" + стр(н) + ".csv", csv.записать(т))
    пусть т2 = csv.разобрать_с_разделителем(стр(прочитать_байты("т" + стр(н) + ".csv")), ";") иначе [["ошибка"]]
    записать("итог_т" + стр(н) + ".csv", csv.записать_для_excel(т2))
}
'''
            код, вывод, ош = выполнить(п, программа)
            self.assertEqual(код, 0, ош)
            for н, т in enumerate(таблицы):
                with self.subTest(н=н):
                    with open(os.path.join(п, f"итог_з{н}.csv"), encoding="utf-8", newline="") as ф:
                        self.assertEqual(list(csv.reader(ф)), т)
                    with open(os.path.join(п, f"итог_т{н}.csv"), encoding="utf-8-sig", newline="") as ф:
                        self.assertEqual(list(csv.reader(ф, delimiter=";")), т)

    def test_разделитель_и_ошибки(self):
        with tempfile.TemporaryDirectory() as п:
            код, вывод, ош = выполнить(п, r'''подключить "csv"
вывод(csv.разобрать("а;б;в\n1;2;3") иначе [])
вывод(csv.разобрать("а\tб\n1\t2") иначе [])
вывод(csv.разобрать("а,\"б;в\",г") иначе [])
вывод(csv.разобрать("") иначе [["ошибка"]])
если ошибка о = csv.разобрать("а,\"б\nв") { вывод(о) }
если ошибка о = csv.разобрать("\"а\"б,в") { вывод(о) }
если ошибка о = csv.разобрать_с_разделителем("а", "\"") { вывод(о) }
''')
        self.assertEqual(код, 0, ош)
        self.assertEqual(вывод.splitlines(), [
            '[["а", "б", "в"], ["1", "2", "3"]]', '[["а", "б"], ["1", "2"]]', '[["а", "б;в", "г"]]', "[]",
            "CSV, строка 1: кавычка не закрыта",
            "CSV, строка 1: после закрывающей кавычки должен идти разделитель",
            "разделитель CSV - один знак, кроме кавычки и перевода строки"])


if __name__ == "__main__":
    unittest.main()
